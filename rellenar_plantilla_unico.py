"""Relleno de plantilla con 1 archivo por RUT (pagos AFP/Isapre/Mutual/Caja).

La plantilla de pagos NO es igual a liquidaciones:
  A=nombre_de_archivo, B=RUT, C=Nombre, ...  (sin documento_asociado_*)
Por eso las columnas se resuelven por encabezado, no por índice fijo.
"""

from __future__ import annotations

import logging
from datetime import date
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from flows.plantilla_schema import HOJA_DOCUMENTOS
from gcs_docs import DocGCS, buscar_docs_tipo, rut_cuerpo

logger = logging.getLogger(__name__)


@dataclass
class ResultadoRellenoUnico:
    plantilla: Path
    periodo_gcs: str
    filas_originales: int = 0
    filas: int = 0
    eliminadas: int = 0
    ruts_eliminados: list[str] = field(default_factory=list)
    docs: list[DocGCS] = field(default_factory=list)


def _mapa_columnas(ws: Worksheet) -> dict[str, int]:
    """header_normalizado → índice 1-based."""
    out: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        raw = ws.cell(1, c).value
        if raw is None:
            continue
        key = str(raw).strip().lower()
        out[key] = c
    return out


def _col(mapa: dict[str, int], *nombres: str) -> int:
    for n in nombres:
        if n.lower() in mapa:
            return mapa[n.lower()]
    raise KeyError(f"No encontré columna {nombres} en encabezados {list(mapa)}")


def rellenar_plantilla_un_documento(
    plantilla: Path,
    *,
    tipo_gcs: str,
    layout: str = "yyyy_mm",
    periodo_gcs: str | None = None,
    guardar_como: Path | None = None,
    referencia: date | None = None,
) -> ResultadoRellenoUnico:
    """
    Escribe nombre del PDF en 'nombre_de_archivo'.
    Solo deja filas con documento en GCS; borra el resto.
    No toca columnas de RUT/Nombre ni inventa documento_asociado_*.
    """
    docs_idx, periodo = buscar_docs_tipo(
        tipo_gcs, periodo=periodo_gcs, layout=layout, referencia=referencia
    )
    wb = load_workbook(plantilla)
    if HOJA_DOCUMENTOS not in wb.sheetnames:
        raise ValueError(f"No existe hoja '{HOJA_DOCUMENTOS}' en {plantilla}")
    ws = wb[HOJA_DOCUMENTOS]
    cols = _mapa_columnas(ws)
    col_archivo = _col(cols, "nombre_de_archivo")
    col_rut = _col(cols, "rut", "RUT")
    # opcionales: limpiar asociados si existen (plantilla estilo liquidaciones)
    cols_asociados = [
        cols[k]
        for k in cols
        if k.startswith("documento_asociado")
    ]

    logger.info(
        "Plantilla columnas: archivo=%s rut=%s asociados=%s headers=%s",
        col_archivo,
        col_rut,
        cols_asociados,
        list(cols.keys()),
    )

    res = ResultadoRellenoUnico(plantilla=plantilla, periodo_gcs=periodo)
    filas_a_borrar: list[int] = []
    a_mantener: list[tuple[int, DocGCS]] = []

    for row in range(2, ws.max_row + 1):
        rut_raw = ws.cell(row, col_rut).value
        if rut_raw is None or str(rut_raw).strip() == "":
            filas_a_borrar.append(row)
            continue
        res.filas_originales += 1
        rut_str = str(rut_raw).strip()
        cuerpo = rut_cuerpo(rut_str)
        doc = docs_idx.get(cuerpo)
        if not doc:
            filas_a_borrar.append(row)
            res.ruts_eliminados.append(rut_str)
            continue
        a_mantener.append((row, doc))

    vistos: set[str] = set()
    for row, doc in a_mantener:
        ws.cell(row, col_archivo, doc.nombre)
        for col in cols_asociados:
            ws.cell(row, col, "")
        if doc.nombre not in vistos:
            res.docs.append(doc)
            vistos.add(doc.nombre)

    for row in sorted(filas_a_borrar, reverse=True):
        ws.delete_rows(row, 1)

    res.eliminadas = len(res.ruts_eliminados)
    res.filas = len(a_mantener)

    dest = guardar_como or plantilla
    wb.save(dest)
    res.plantilla = dest

    logger.info(
        "Plantilla 1-doc tipo=%s periodo=%s | orig=%s keep=%s elim=%s → %s",
        tipo_gcs,
        periodo,
        res.filas_originales,
        res.filas,
        res.eliminadas,
        dest,
    )
    if res.ruts_eliminados:
        logger.warning(
            "RUTs sin documento (%s): %s",
            len(res.ruts_eliminados),
            res.ruts_eliminados[:15],
        )
    if res.filas == 0:
        raise ValueError(
            f"Ningún RUT tiene '{tipo_gcs}' en GCS para periodo '{periodo}'."
        )
    return res
