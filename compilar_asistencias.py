"""
Mapea RUT→faena desde plantilla de liquidaciones y compila
un PDF por faena (periodo mes anterior en GCS).

Usado por Libro de Asistencia (tipo GCS 'Asistencia') y por las cargas agrupadas
por faena de Pagos Mutualidades / Cajas de Compensación (tipo GCS 'Cotizaciones').
"""

from __future__ import annotations

import json
import logging
import re
import unicodedata
from collections import defaultdict
from datetime import date
from dataclasses import dataclass, field
from pathlib import Path

import fitz  # pymupdf
import openpyxl

from browser import MESES_ES, mes_anterior
from config import DOWNLOAD_DIR
from gcs_docs import (
    DocGCS,
    buscar_docs_tipo,
    descargar_doc,
    periodo_gcs_texto,
    periodo_gcs_yyyy_mm,
    rut_cuerpo,
)

logger = logging.getLogger(__name__)

HOJA = "Documentos"


@dataclass(frozen=True)
class FuenteDocsFaena:
    """
    De dónde salen los PDF por RUT que se fusionan por faena (parámetro por flujo).

    - tipo_gcs / layout definen la ruta GCS (ver gcs_docs._globs_periodo):
        mes_anio → Trabajadores/{rut}/{tipo}/{mes año}/*.pdf  (+ legacy **/{tipo}/{mes año}/)
        yyyy_mm  → Trabajadores/{rut}/{tipo}/{YYYY}/{MM}/*.pdf (+ legacy **/{tipo}/{YYYY}/{MM}/)
    - RUT del PDF: prefijo del nombre `{rut}[-dv]_...` o carpeta `Trabajadores/{rut}/`
      (gcs_docs._rut_cuerpo_de_blob). 1 PDF por RUT (gcs_docs.indice_por_rut).
    - etiqueta: nombre del fusionado `{faena_slug}_{etiqueta}_{periodo}.pdf`.
    - orden: 'excel' (aparición en la base) o 'rut' (numérico).
    """

    tipo_gcs: str
    layout: str
    etiqueta: str
    orden: str = "excel"
    staging_nombre: str = "staging_asistencias"

    def patron_gcs(self, periodo: str, anio: str = "{YYYY}", mes: str = "{MM}") -> str:
        if self.layout == "yyyy_mm":
            return f"Trabajadores/{{rut}}/{self.tipo_gcs}/{anio}/{mes}/*.pdf"
        return f"Trabajadores/{{rut}}/{self.tipo_gcs}/{periodo}/*.pdf"


FUENTE_ASISTENCIA = FuenteDocsFaena(
    tipo_gcs="Asistencia", layout="mes_anio", etiqueta="Asistencia", orden="excel"
)


@dataclass
class ResultadoCompilacion:
    periodo: str
    por_faena: dict[str, list[str]]  # faena → ruts cuerpo
    pdfs: dict[str, Path] = field(default_factory=dict)
    sin_pdf: dict[str, list[str]] = field(default_factory=dict)
    mapa_path: Path | None = None
    staging: Path | None = None


def _slug_faena(nombre: str) -> str:
    s = unicodedata.normalize("NFKD", nombre)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^\w\-]+", "_", s, flags=re.UNICODE).strip("_")
    return s[:80] or "sin_faena"


def leer_ruts_por_faena(plantilla: Path) -> dict[str, list[str]]:
    """
    Lee columnas RUT y faena de la hoja Documentos.
    Mantiene orden de aparición; un RUT puede figurar en varias faenas.
    """
    wb = openpyxl.load_workbook(plantilla, data_only=True, read_only=True)
    if HOJA not in wb.sheetnames:
        raise ValueError(f"No existe hoja '{HOJA}' en {plantilla}")
    ws = wb[HOJA]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows)
    try:
        i_rut = headers.index("RUT")
        i_faena = headers.index("faena")
    except ValueError as e:
        raise ValueError(f"Faltan columnas RUT/faena en {plantilla}: {headers}") from e

    por_faena: dict[str, list[str]] = defaultdict(list)
    vistos: set[tuple[str, str]] = set()
    for row in rows:
        if not row or i_rut >= len(row):
            continue
        rut_raw = row[i_rut]
        if rut_raw is None or str(rut_raw).strip() == "":
            continue
        cuerpo = rut_cuerpo(str(rut_raw))
        if not cuerpo:
            continue
        faena_raw = row[i_faena] if i_faena < len(row) else None
        faena = str(faena_raw).strip() if faena_raw else "(sin_faena)"
        key = (cuerpo, faena)
        if key in vistos:
            continue
        vistos.add(key)
        por_faena[faena].append(cuerpo)

    wb.close()
    return dict(sorted(por_faena.items(), key=lambda kv: (-len(kv[1]), kv[0])))


def fusionar_pdfs(paths: list[Path], dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    out = fitz.open()
    try:
        for p in paths:
            with fitz.open(p) as src:
                out.insert_pdf(src)
        out.save(str(dest))
    finally:
        out.close()
    return dest


def compilar_asistencias_por_faena(
    plantilla: Path,
    *,
    periodo: str | None = None,
    staging: Path | None = None,
) -> ResultadoCompilacion:
    """Libro de Asistencia: 1 PDF por faena con las asistencias de sus RUT."""
    return compilar_docs_por_faena(
        plantilla, FUENTE_ASISTENCIA, periodo=periodo, staging=staging
    )


def compilar_docs_por_faena(
    plantilla: Path,
    fuente: FuenteDocsFaena,
    *,
    periodo: str | None = None,
    staging: Path | None = None,
    referencia: date | None = None,
) -> ResultadoCompilacion:
    """
    Agrupa los RUT de la plantilla de liquidaciones (base) por columna 'faena' y
    fusiona el PDF GCS de cada RUT (según `fuente`) en un único PDF por faena:
    `{slug_faena}_{fuente.etiqueta}_{periodo}.pdf` en `staging/por_faena`.
    Los RUT sin PDF quedan en resultado.sin_pdf[faena] (log + resumen JSON).
    """
    plantilla = Path(plantilla)
    periodo = periodo or periodo_gcs_texto()
    por_faena = leer_ruts_por_faena(plantilla)

    idx, periodo_res = buscar_docs_tipo(
        fuente.tipo_gcs, periodo=periodo, layout=fuente.layout, referencia=referencia
    )
    if periodo_res != periodo:
        logger.info("Periodo GCS resuelto: %s", periodo_res)
        periodo = periodo_res
    if fuente.layout == "yyyy_mm":
        _, _anio, _mes = periodo_gcs_yyyy_mm(referencia)
        patron = fuente.patron_gcs(periodo, _anio, _mes)
    else:
        patron = fuente.patron_gcs(periodo)
    logger.info(
        "Fuente por faena: tipo=%s ruta=%s | RUT=prefijo nombre o carpeta | "
        "%s RUT con PDF en GCS | orden=%s | salida=*_%s_%s.pdf",
        fuente.tipo_gcs, patron, len(idx), fuente.orden, fuente.etiqueta, periodo,
    )
    etiqueta = fuente.etiqueta
    orden = fuente.orden
    staging_nombre = fuente.staging_nombre

    staging = Path(
        staging
        or DOWNLOAD_DIR / staging_nombre / periodo.replace(" ", "_")
    )
    staging.mkdir(parents=True, exist_ok=True)
    raw_dir = staging / "_raw"
    out_dir = staging / "por_faena"
    out_dir.mkdir(parents=True, exist_ok=True)

    mapa = {
        "periodo": periodo,
        "tipo_gcs": fuente.tipo_gcs,
        "ruta_gcs": patron,
        "orden": fuente.orden,
        "plantilla": str(plantilla.resolve()),
        "total_faenas": len(por_faena),
        "total_ruts": sum(len(v) for v in por_faena.values()),
        "faenas": {
            f: {"ruts": ruts, "n": len(ruts)} for f, ruts in por_faena.items()
        },
    }
    mapa_path = staging / "mapa_ruts_por_faena.json"
    mapa_path.write_text(
        json.dumps(mapa, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    logger.info(
        "Mapa RUT/faena: %s faenas, %s pares → %s",
        mapa["total_faenas"],
        mapa["total_ruts"],
        mapa_path,
    )

    resultado = ResultadoCompilacion(
        periodo=periodo,
        por_faena=por_faena,
        mapa_path=mapa_path,
        staging=staging,
    )

    for faena, ruts in por_faena.items():
        if orden == "rut":
            ruts = sorted(ruts, key=lambda r: (len(r), r))
        docs: list[DocGCS] = []
        faltan: list[str] = []
        for cuerpo in ruts:
            doc = idx.get(cuerpo)
            if doc:
                docs.append(doc)
            else:
                faltan.append(cuerpo)
        if faltan:
            resultado.sin_pdf[faena] = faltan
            logger.warning(
                "[%s] %s RUT(s) sin PDF: %s",
                faena,
                len(faltan),
                ", ".join(faltan[:8]) + ("…" if len(faltan) > 8 else ""),
            )

        if not docs:
            logger.warning("[%s] sin PDFs — no se genera compilado", faena)
            continue

        # Subcarpeta por RUT: evita colisiones si dos RUT tienen el mismo nombre de archivo
        paths = [descargar_doc(d, raw_dir / (d.rut_cuerpo or "_sin_rut")) for d in docs]
        slug = _slug_faena(faena)
        dest = out_dir / f"{slug}_{etiqueta}_{periodo}.pdf"
        fusionar_pdfs(paths, dest)
        resultado.pdfs[faena] = dest
        logger.info(
            "[%s] compilado %s PDFs → %s (%s bytes)",
            faena,
            len(paths),
            dest.name,
            dest.stat().st_size,
        )

    resumen = {
        "periodo": periodo,
        "compilados": len(resultado.pdfs),
        "faenas_sin_pdf": len(resultado.sin_pdf),
        "ruts_sin_pdf": sum(len(v) for v in resultado.sin_pdf.values()),
        "pdfs": {f: str(p) for f, p in resultado.pdfs.items()},
        "sin_pdf": resultado.sin_pdf,
    }
    (staging / "resumen_compilacion.json").write_text(
        json.dumps(resumen, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return resultado


def main() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
    )
    import argparse
    import sys

    parser = argparse.ArgumentParser(description="Compilar asistencias por faena")
    parser.add_argument(
        "plantilla",
        nargs="?",
        help="Excel plantilla liquidaciones (cmsd_plantilla_*.xlsx)",
    )
    parser.add_argument(
        "--periodo",
        default=None,
        help="Periodo GCS, ej: 'julio 2026' (default: mes anterior)",
    )
    args = parser.parse_args()

    if args.plantilla:
        plantilla = Path(args.plantilla)
    else:
        candidatos = sorted(
            DOWNLOAD_DIR.glob("cmsd_plantilla_*.xlsx"),
            key=lambda p: p.stat().st_mtime,
            reverse=True,
        )
        if not candidatos:
            print("No hay cmsd_plantilla_*.xlsx en downloads/", file=sys.stderr)
            sys.exit(1)
        plantilla = candidatos[0]

    p = args.periodo
    if not p:
        m = mes_anterior()
        p = f"{MESES_ES[m.month]} {m.year}"

    r = compilar_asistencias_por_faena(plantilla, periodo=p)
    print(
        f"OK: {len(r.pdfs)} PDF(s) en {r.staging / 'por_faena'} | "
        f"sin PDF: {sum(len(v) for v in r.sin_pdf.values())} RUT(s) | "
        f"mapa: {r.mapa_path}"
    )


if __name__ == "__main__":
    main()
