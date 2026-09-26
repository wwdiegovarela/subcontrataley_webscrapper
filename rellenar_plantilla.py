"""Rellena la plantilla Excel de liquidaciones con nombres desde GCS."""

from __future__ import annotations

import logging
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from flows.plantilla_schema import HOJA_DOCUMENTOS
from generar_transferencia import crear_transferencia_una_linea, liquidacion_de_finiquitado
from gcs_docs import (
    DocGCS,
    buscar_liquidaciones_y_transferencias,
    descargar_doc,
    rut_cuerpo,
    tiene_finiquito,
)
from parse_liquidacion import parse_total_liquido

logger = logging.getLogger(__name__)

COL_NOMBRE_ARCHIVO = 1  # A
COL_DOC_ASOCIADO_1 = 2  # B
COL_DOC_ASOCIADO_2 = 3  # C
COL_DOC_ASOCIADO_3 = 4  # D
COL_DOC_ASOCIADO_4 = 5  # E
COL_DOC_ASOCIADO_5 = 6  # F
COL_RUT = 7  # G

COLS_DOC_ASOCIADO_VACIAS = (
    COL_DOC_ASOCIADO_2,
    COL_DOC_ASOCIADO_3,
    COL_DOC_ASOCIADO_4,
    COL_DOC_ASOCIADO_5,
)


@dataclass
class ResultadoRelleno:
    plantilla: Path
    periodo_gcs: str
    filas_originales: int = 0
    filas: int = 0  # filas que quedan (completas o líquido 0)
    eliminadas: int = 0
    liquido_cero: int = 0
    ruts_eliminados: list[str] = field(default_factory=list)
    docs_liquidacion: list[DocGCS] = field(default_factory=list)
    docs_transferencia: list[DocGCS] = field(default_factory=list)
    transferencias_generadas: list[Path] = field(default_factory=list)

    @property
    def con_liquidacion(self) -> int:
        return self.filas

    @property
    def con_transferencia(self) -> int:
        return len(self.docs_transferencia)

    @property
    def sin_liquidacion(self) -> list[str]:
        return self.ruts_eliminados

    @property
    def sin_transferencia(self) -> list[str]:
        return self.ruts_eliminados


def _es_liquido_cero(liq: DocGCS, cache_dir: Path) -> bool:
    """Descarga temporal del PDF y evalúa TOTAL LIQUIDO == 0."""
    try:
        local = descargar_doc(liq, cache_dir)
        monto = parse_total_liquido(local)
        return monto == 0
    except Exception as exc:
        logger.warning("No pude evaluar líquido de %s: %s", liq.nombre, exc)
        return False


def _nombre_transferencia(tr: DocGCS | Path | None) -> str:
    if tr is None:
        return ""
    if isinstance(tr, Path):
        return tr.name
    return tr.nombre


def rellenar_plantilla_liquidaciones(
    plantilla: Path,
    *,
    periodo_gcs: str | None = None,
    guardar_como: Path | None = None,
    solo_completos: bool = True,
    dir_transferencias_generadas: Path | None = None,
) -> ResultadoRelleno:
    """
    Columna A = Liquidación, B = Transferencia (nombres GCS o generados).

    Reglas (solo_completos=True):
    - RUT con liquidación + transferencia → se mantiene
    - RUT con liquidación, sin transferencia, TOTAL LIQUIDO == 0 → se mantiene (B="")
    - RUT finiquitado, con liquidación, sin transferencia y TOTAL LIQUIDO > 0
      → se crea el comprobante de una línea y se pone en B
    - Resto → se borra del Excel
    """
    liqs, trs, periodo = buscar_liquidaciones_y_transferencias(periodo=periodo_gcs)
    dir_gen = Path(
        dir_transferencias_generadas
        or (plantilla.parent / "transferencias_generadas")
    )
    wb = load_workbook(plantilla)
    if HOJA_DOCUMENTOS not in wb.sheetnames:
        raise ValueError(f"No existe hoja '{HOJA_DOCUMENTOS}' en {plantilla}")
    ws = wb[HOJA_DOCUMENTOS]

    res = ResultadoRelleno(plantilla=plantilla, periodo_gcs=periodo)
    filas_a_borrar: list[int] = []
    # (row, rut, liq, tr|None)
    a_mantener: list[tuple[int, str, DocGCS, DocGCS | None]] = []

    with tempfile.TemporaryDirectory(prefix="liq_parse_") as tmp:
        cache_dir = Path(tmp)

        for row in range(2, ws.max_row + 1):
            rut_raw = ws.cell(row, COL_RUT).value
            if rut_raw is None or str(rut_raw).strip() == "":
                filas_a_borrar.append(row)
                continue

            res.filas_originales += 1
            rut_str = str(rut_raw).strip()
            cuerpo = rut_cuerpo(rut_str)
            liq = liqs.get(cuerpo)
            tr = trs.get(cuerpo)

            if not solo_completos:
                if liq or tr:
                    a_mantener.append((row, rut_str, liq, tr))  # type: ignore[arg-type]
                else:
                    filas_a_borrar.append(row)
                    res.ruts_eliminados.append(rut_str)
                continue

            # Par completo
            if liq and tr:
                a_mantener.append((row, rut_str, liq, tr))
                continue

            # Solo liquidación: líquido 0 queda sin transferencia.
            # El comprobante de una línea solo si está finiquitado y el líquido es > 0.
            if liq and not tr:
                local_liq = descargar_doc(liq, cache_dir)
                monto = parse_total_liquido(local_liq)
                if monto == 0:
                    logger.info(
                        "RUT %s: líquido=0 → liquidación sin transferencia",
                        rut_str,
                    )
                    a_mantener.append((row, rut_str, liq, None))
                    res.liquido_cero += 1
                    continue
                finiquitado = liquidacion_de_finiquitado(local_liq) or tiene_finiquito(
                    cuerpo
                )
                generada = None
                if finiquitado and monto is not None and monto > 0:
                    generada = crear_transferencia_una_linea(
                        local_liq, dir_gen, periodo, monto
                    )
                if generada is None:
                    filas_a_borrar.append(row)
                    res.ruts_eliminados.append(rut_str)
                    logger.info(
                        "RUT %s: sin transferencia (finiquitado=%s)",
                        rut_str,
                        finiquitado,
                    )
                    continue
                logger.info("RUT %s: transferencia creada %s", rut_str, generada.name)
                a_mantener.append((row, rut_str, liq, generada))
                continue

            # Sin liquidación (con o sin transferencia)
            filas_a_borrar.append(row)
            res.ruts_eliminados.append(rut_str)

    vistos_liq: set[str] = set()
    vistos_tr: set[str] = set()
    for row, _rut, liq, tr in a_mantener:
        if not liq:
            filas_a_borrar.append(row)
            continue
        ws.cell(row, COL_NOMBRE_ARCHIVO, liq.nombre)
        ws.cell(row, COL_DOC_ASOCIADO_1, _nombre_transferencia(tr))
        for col in COLS_DOC_ASOCIADO_VACIAS:
            ws.cell(row, col, "")
        if liq.nombre not in vistos_liq:
            res.docs_liquidacion.append(liq)
            vistos_liq.add(liq.nombre)
        if isinstance(tr, Path):
            if tr not in res.transferencias_generadas:
                res.transferencias_generadas.append(tr)
        elif tr and tr.nombre not in vistos_tr:
            res.docs_transferencia.append(tr)
            vistos_tr.add(tr.nombre)

    for row in sorted(set(filas_a_borrar), reverse=True):
        ws.delete_rows(row, 1)

    res.eliminadas = len(res.ruts_eliminados)
    res.filas = len([x for x in a_mantener if x[2] is not None])

    dest = guardar_como or plantilla
    wb.save(dest)
    res.plantilla = dest

    logger.info(
        "Plantilla periodo=%s | orig=%s | keep=%s (liq0=%s, tr_generadas=%s) | elim=%s → %s",
        periodo,
        res.filas_originales,
        res.filas,
        res.liquido_cero,
        len(res.transferencias_generadas),
        res.eliminadas,
        dest,
    )
    if res.ruts_eliminados:
        logger.warning(
            "RUTs eliminados (%s): %s",
            len(res.ruts_eliminados),
            res.ruts_eliminados[:15],
        )
    if res.filas == 0:
        raise ValueError(
            f"Ningún RUT válido para periodo '{periodo}' "
            "(par completo o liquidación con líquido 0)."
        )
    return res
