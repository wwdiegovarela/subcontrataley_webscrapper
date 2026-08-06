"""Leer plantilla Excel de trabajadores descargada de Subcontrataley."""

from __future__ import annotations

import logging
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from gcs_docs import rut_cuerpo

logger = logging.getLogger(__name__)

HOJA_TRABAJADORES = "Trabajadores"


@dataclass
class TrabajadorSCL:
    contrato_id: str
    fc_id: str
    rut_raw: str
    rut_cuerpo: str
    nombres: str
    tipo_contrato: str
    razon_social: str
    faena: str
    servicio: str
    ingreso_faena: str
    fecha_termino: str
    fecha_desvinculacion: str
    tipos_movimiento_id: str

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _cell(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def leer_listado_trabajadores(path: Path | str) -> list[TrabajadorSCL]:
    """
    Plantilla finiquito / trabajadores SCL (hoja Trabajadores).

    Columnas: contrato_id, fc_id, rut, nombres, tipo_contrato, razon_social,
    faena, servicio, ingreso_faena, fecha_termino, fecha_desvinculacion,
    tipos_movimiento_id.
    """
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    if HOJA_TRABAJADORES not in wb.sheetnames:
        wb.close()
        raise ValueError(
            f"Hoja '{HOJA_TRABAJADORES}' no está en {path.name}: {wb.sheetnames}"
        )
    ws = wb[HOJA_TRABAJADORES]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        wb.close()
        return []

    cols = {str(h).strip().lower(): i for i, h in enumerate(header) if h is not None}
    for req in ("rut", "nombres", "faena"):
        if req not in cols:
            wb.close()
            raise ValueError(f"Falta columna '{req}' en {path}. Header: {header}")

    def g(row: tuple, key: str) -> str:
        i = cols.get(key)
        if i is None or i >= len(row):
            return ""
        return _cell(row[i])

    out: list[TrabajadorSCL] = []
    vistos: set[str] = set()
    for row in rows:
        if not row:
            continue
        raw = g(row, "rut")
        if not raw:
            continue
        cuerpo = rut_cuerpo(raw)
        if not cuerpo or cuerpo in vistos:
            continue
        vistos.add(cuerpo)
        out.append(
            TrabajadorSCL(
                contrato_id=g(row, "contrato_id"),
                fc_id=g(row, "fc_id"),
                rut_raw=raw,
                rut_cuerpo=cuerpo,
                nombres=g(row, "nombres"),
                tipo_contrato=g(row, "tipo_contrato"),
                razon_social=g(row, "razon_social"),
                faena=g(row, "faena"),
                servicio=g(row, "servicio"),
                ingreso_faena=g(row, "ingreso_faena"),
                fecha_termino=g(row, "fecha_termino"),
                fecha_desvinculacion=g(row, "fecha_desvinculacion"),
                tipos_movimiento_id=g(row, "tipos_movimiento_id"),
            )
        )
    wb.close()
    logger.info("SCL listado %s: %s RUTs únicos", path.name, len(out))
    return out
