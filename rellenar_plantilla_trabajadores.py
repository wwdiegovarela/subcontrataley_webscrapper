"""
Rellena plantilla Subcontrataley de ingreso de trabajadores (hoja Trabajadores).

Fuentes:
  - BigQuery cr_reclutamiento_empleados (dirección, comuna, celular, jornada, fechas, instalación, RUT-DV)
  - BigQuery cr_asistencia_hist_tb (fecha_ingreso_faena = 1ª marca asistencia)
  - ControlRoll token contratos (nombres, AFP, salud, sexo, nacionalidad, etc.)
  - DICCIONARIO_INSTALACIONES_WALMART (instalación → faena SCL)
  - Catálogos embebidos en la plantilla Excel

Uso:
  python rellenar_plantilla_trabajadores.py \\
    --plantilla "C:/Users/Diego/Downloads/cmct_g_20260805095021_43255_167_0_18775_0.xlsx" \\
    --ruts 11860305,14182089,...
"""

from __future__ import annotations

import argparse
import logging
import re
import shutil
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from bq_asistencia import consultar_primera_asistencia, corte_ingreso_faena
from bq_empleados import consultar_empleados_activos
from config import DOWNLOAD_DIR
from cr_contratos import TOKEN_CR_CONTRATOS, _consulta_cr_json
from gcs_docs import rut_cuerpo

logger = logging.getLogger(__name__)

# --- Mapeos fijos acordados ---
CARGO_ID_DEFAULT = 553
DUENNO_DEFAULT = "NO"

AFP_MAP = {
    "MODELO": 7,
    "PROVIDA": 5,
    "UNO": 34,
    "HABITAT": 3,
    "CAPITAL": 1,
    "PLANVITAL": 4,
    "PLAN VITAL": 4,
    "CUPRUM": 2,
    "INP SSS": 6,
    "IPS (INP)": 6,
}

SALUD_MAP = {
    "FONASA": 19,
    "CRUZ BLANCA": 22,
    "NUEVA MASVIDA": 10,
    "BANMEDICA": 7,
    "BANMÉDICA": 7,
    "CONSALUD": 8,
    "COLMENA": 1,
    "VIDA TRES": 15,
    "OPTIMA": 21,
    "DIPRECA": 18,
    "CAPREDENA": 17,
}

JORNADA_MAP = {
    "4X4": 17,
    "6X2": 22,
    "6X1": 25,
    "DIA": 1,
    "DÍA": 1,
}

SEXO_MAP = {
    "HOMBRE": 1,
    "MASCULINO": 1,
    "MUJER": 2,
    "FEMENINO": 2,
}

REGIMENES_PENSIONADO = {5, 7, 8}

HEADERS = [
    "contrato_id",
    "pais_id",
    "documento_id",
    "rut",
    "nombres",
    "apellido_paterno",
    "apellido_materno",
    "genero_id",
    "fecha_nacimiento",
    "comuna_id",
    "direccion",
    "telefono",
    "isapre_id",
    "afp_id",
    "pensionado",
    "duenno",
    "email",
    "cargo_id",
    "articulo_22",
    "contratotipo_id",
    "fecha_inicio_contrato",
    "fecha_fin_contrato",
    "contratojornada_id",
    "fecha_ingreso_faena",
    "servicio_id",
    "faena_id",
]


def _norm(s: Any) -> str:
    t = unicodedata.normalize("NFKD", str(s or ""))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", t.upper().replace("-", " ")).strip()


def _regimen_code(s: Any) -> int | None:
    m = re.match(r"^(\d+)", _norm(s))
    return int(m.group(1)) if m else None


def _load_walmart_dict() -> dict[str, str]:
    from mapa_instalaciones_walmart import DICCIONARIO_INSTALACIONES_WALMART

    out: dict[str, str] = {}
    for r in DICCIONARIO_INSTALACIONES_WALMART:
        key = _norm(r["instalacion"])
        scl = r["instalacion_subcontrataley"]
        if key in out and "(" not in scl and "(" in out[key]:
            continue
        if key in out and "(" in scl and "(" not in out[key]:
            out[key] = scl
            continue
        out[key] = scl
    return out


def _catalogos_desde_plantilla(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=True)
    sheets = {s: wb[s] for s in wb.sheetnames}

    def sheet_by_part(part: str):
        for s in wb.sheetnames:
            if part.lower() in s.lower():
                return wb[s]
        raise KeyError(part)

    comunas: dict[str, int] = {}
    for r in sheet_by_part("Comunas").iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        comunas[_norm(r[1])] = int(r[0])

    pais_por_nac: dict[str, int] = {}
    pais_por_nombre: dict[str, int] = {}
    for r in sheet_by_part("Pa").iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        pais_por_nombre[_norm(r[1])] = int(r[0])
        if len(r) > 2 and r[2]:
            pais_por_nac[_norm(r[2])] = int(r[0])

    faenas: dict[str, int] = {}
    for r in sheets["Faenas"].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        faenas[_norm(r[1])] = int(r[0])

    servicios: dict[str, int] = {}
    for r in sheets["Servicios"].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        servicios[_norm(r[4])] = int(r[0])

    jornadas: dict[str, int] = {}
    for r in sheets["Jornada"].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        jornadas[_norm(r[1])] = int(r[0])

    wb.close()
    return {
        "comunas": comunas,
        "pais_por_nac": pais_por_nac,
        "pais_por_nombre": pais_por_nombre,
        "faenas": faenas,
        "servicios": servicios,
        "jornadas": jornadas,
    }


def _parse_fecha_cr(s: Any) -> str:
    """CR '24-11-1971' o datetime → 'YYYY-MM-DD'."""
    if s is None or str(s).strip() == "":
        return ""
    if isinstance(s, (date, datetime)):
        return (s.date() if isinstance(s, datetime) else s).isoformat()
    t = str(s).strip()
    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(t[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return t


def _map_afp(afp: str, pensionado: str, regimen: str) -> int | None:
    a = _norm(afp)
    if a == "SIN AFP":
        code = _regimen_code(regimen)
        if _norm(pensionado) == "SI" or code in REGIMENES_PENSIONADO:
            return 10
        return None  # excluir
    # compact PLANVITAL
    a2 = a.replace(" ", "")
    if a in AFP_MAP:
        return AFP_MAP[a]
    if a2 in AFP_MAP:
        return AFP_MAP[a2]
    if "CAPITAL" in a:
        return 1
    if "PLAN" in a and "VITAL" in a:
        return 4
    if "INP" in a or "IPS" in a:
        return 6
    return None


def _map_contratotipo(tipo: str) -> int:
    t = _norm(tipo)
    if "POSTERIOR INDEFINIDO" in t or t == "PLAZO INDEFINIDO" or "INDEFINIDO" in t:
        return 2
    if "PLAZO FIJO" in t:
        return 1
    return 2


def _map_jornada(jornada: str, cat: dict[str, int]) -> int | None:
    j = _norm(jornada).replace(" ", "")
    if j in JORNADA_MAP:
        return JORNADA_MAP[j]
    # try catalog "4 X 4"
    spaced = _norm(jornada)
    if spaced in cat:
        return cat[spaced]
    # fuzzy 4 X 4
    for k, v in cat.items():
        if k.replace(" ", "") == j:
            return v
    return None


def _map_comuna(comuna: str, cat: dict[str, int]) -> int | None:
    n = _norm(comuna)
    if n in cat:
        return cat[n]
    # soft: startswith / contains
    for k, v in cat.items():
        if k == n or k.startswith(n) or n.startswith(k):
            return v
    return None


def _map_pais(nacionalidad: str, cat_nac: dict[str, int], cat_pais: dict[str, int]) -> int:
    n = _norm(nacionalidad)
    if n in cat_nac:
        return cat_nac[n]
    # CHILENA / CHILE
    if "CHILEN" in n:
        return 1
    if n in cat_pais:
        return cat_pais[n]
    return 1


def _map_salud(salud: str) -> int:
    s = _norm(salud)
    if s in SALUD_MAP:
        return SALUD_MAP[s]
    s2 = s.replace(" ", "")
    for k, v in SALUD_MAP.items():
        if _norm(k).replace(" ", "") == s2:
            return v
    if "FONASA" in s:
        return 19
    return 19


def _fetch_cr_rows_by_rut(token: str | None = None) -> dict[str, dict]:
    """Todas las filas del reporte CR indexadas por rut_cuerpo (última gana)."""
    rows = _consulta_cr_json(token or TOKEN_CR_CONTRATOS)
    out: dict[str, dict] = {}
    for row in rows:
        cuerpo = rut_cuerpo(row.get("RUT") or "")
        if cuerpo:
            out[cuerpo] = row
    return out


def _si_no(v: Any, default: str = "NO") -> str:
    n = _norm(v)
    if n in ("SI", "SÍ", "YES", "TRUE", "1"):
        return "SI"
    if n in ("NO", "FALSE", "0"):
        return "NO"
    return default


def construir_filas(
    *,
    plantilla: Path,
    ruts: list[str],
) -> tuple[list[dict[str, Any]], list[str]]:
    cats = _catalogos_desde_plantilla(plantilla)
    walmart = _load_walmart_dict()
    cr_by_rut = _fetch_cr_rows_by_rut()
    empleados = {e.rut_cuerpo: e for e in consultar_empleados_activos()}

    filas: list[dict[str, Any]] = []
    excluidos: list[str] = []

    for raw in ruts:
        cuerpo = rut_cuerpo(raw)
        emp = empleados.get(cuerpo)
        cr = cr_by_rut.get(cuerpo)
        if not emp:
            excluidos.append(f"{cuerpo}: no está en BQ activos Walmart")
            continue
        if not cr:
            excluidos.append(f"{cuerpo}: no está en reporte CR contratos")
            continue

        afp_id = _map_afp(cr.get("AFP"), cr.get("PENSIONADO"), cr.get("Régimen de Pensiones"))
        if afp_id is None:
            excluidos.append(
                f"{cuerpo}: SIN AFP sin régimen pensionado ni PENSIONADO=SI (excluido)"
            )
            continue

        scl_faena = walmart.get(_norm(emp.instalacion))
        if not scl_faena:
            excluidos.append(f"{cuerpo}: instalación sin mapeo: {emp.instalacion}")
            continue
        faena_id = cats["faenas"].get(_norm(scl_faena))
        servicio_id = cats["servicios"].get(_norm(scl_faena))
        if not faena_id or not servicio_id:
            excluidos.append(
                f"{cuerpo}: faena SCL '{scl_faena}' no está en catálogo plantilla"
            )
            continue

        filas.append(
            {
                "_emp": emp,
                "_cr": cr,
                "_afp_id": afp_id,
                "_faena_id": faena_id,
                "_servicio_id": servicio_id,
                "_scl_faena": scl_faena,
            }
        )

    # Enrich with full BQ fields (comuna, direccion, celular, jornada)
    from bq_empleados import _cliente, PROJECT_ID, TABLE_EMPLEADOS
    from google.cloud import bigquery

    cuerpos = [f["_emp"].rut_cuerpo for f in filas]
    if not cuerpos:
        return [], excluidos

    primera_asis = consultar_primera_asistencia(ruts=cuerpos)
    corte_faena = corte_ingreso_faena()
    logger.info(
        "Corte ingreso faena: >= %s (excluye 1ª asistencia más antigua)",
        corte_faena.isoformat(),
    )

    client = _cliente()
    job = bigquery.QueryJobConfig(
        query_parameters=[bigquery.ArrayQueryParameter("ruts", "STRING", cuerpos)]
    )
    sql = f"""
    SELECT
      CAST(rut AS STRING) AS rut,
      CAST(rut_dv AS STRING) AS rut_dv,
      instalacion,
      comuna,
      direccion,
      celular,
      jornada,
      fecha_de_ingreso
    FROM `{TABLE_EMPLEADOS}`
    WHERE CAST(rut AS STRING) IN UNNEST(@ruts)
    """
    bq_extra = {
        rut_cuerpo(r["rut"]): dict(r)
        for r in client.query(sql, job_config=job).result()
    }

    out_rows: list[dict[str, Any]] = []
    for item in filas:
        emp = item["_emp"]
        cr = item["_cr"]
        extra = bq_extra.get(emp.rut_cuerpo, {})
        comuna_id = _map_comuna(extra.get("comuna") or "", cats["comunas"])
        if not comuna_id:
            excluidos.append(
                f"{emp.rut_cuerpo}: comuna sin mapeo: {extra.get('comuna')}"
            )
            continue
        jornada_id = _map_jornada(extra.get("jornada") or "", cats["jornadas"])
        if not jornada_id:
            excluidos.append(
                f"{emp.rut_cuerpo}: jornada sin mapeo: {extra.get('jornada')}"
            )
            continue

        rut_fmt = str(extra.get("rut_dv") or emp.rut_dv or emp.rut_cuerpo).strip()
        fecha_ing = _parse_fecha_cr(extra.get("fecha_de_ingreso"))
        primera = primera_asis.get(emp.rut_cuerpo)
        if not primera:
            excluidos.append(f"{emp.rut_cuerpo}: sin primera marca de asistencia en BQ")
            continue
        if primera < corte_faena:
            excluidos.append(
                f"{emp.rut_cuerpo}: ingreso_faena {primera.isoformat()} "
                f"anterior a corte {corte_faena.isoformat()} (>1 mes)"
            )
            continue
        fecha_faena = primera.isoformat()
        contratotipo = _map_contratotipo(cr.get("TIPO CONTRATO"))
        fecha_fin = ""
        if contratotipo == 1:
            fecha_fin = _parse_fecha_cr(cr.get("Fecha de Término"))

        genero = SEXO_MAP.get(_norm(cr.get("SEXO")))
        if not genero:
            excluidos.append(f"{emp.rut_cuerpo}: sexo sin mapeo: {cr.get('SEXO')}")
            continue

        row = {
            "contrato_id": "",
            "pais_id": _map_pais(
                cr.get("NACIONALIDAD"),
                cats["pais_por_nac"],
                cats["pais_por_nombre"],
            ),
            "documento_id": 1,
            "rut": rut_fmt,
            "nombres": str(cr.get("Nombres") or "").strip(),
            "apellido_paterno": str(cr.get("Ap. Paterno") or "").strip(),
            "apellido_materno": str(cr.get("Ap. Materno") or "").strip(),
            "genero_id": genero,
            "fecha_nacimiento": _parse_fecha_cr(cr.get("FECHA NACIMIENTO")),
            "comuna_id": comuna_id,
            "direccion": str(extra.get("direccion") or "").strip(),
            "telefono": str(extra.get("celular") or "").strip(),
            "isapre_id": _map_salud(cr.get("SALUD")),
            "afp_id": item["_afp_id"],
            "pensionado": _si_no(cr.get("PENSIONADO")),
            "duenno": DUENNO_DEFAULT,
            "email": str(cr.get("Mail") or "").strip(),
            "cargo_id": CARGO_ID_DEFAULT,
            "articulo_22": "NO",
            "contratotipo_id": contratotipo,
            "fecha_inicio_contrato": fecha_ing,
            "fecha_fin_contrato": fecha_fin,
            "contratojornada_id": jornada_id,
            "fecha_ingreso_faena": fecha_faena,
            "servicio_id": item["_servicio_id"],
            "faena_id": item["_faena_id"],
        }
        out_rows.append(row)

    return out_rows, excluidos


def escribir_plantilla(
    plantilla_vacia: Path,
    filas: list[dict[str, Any]],
    destino: Path,
) -> Path:
    destino.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(plantilla_vacia, destino)
    wb = load_workbook(destino)
    ws = wb["Trabajadores"]
    # clear existing data rows (keep header)
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    for i, h in enumerate(HEADERS):
        if i >= len(header) or header[i] != h:
            # still write by HEADER order positions 1..26
            pass
    for fila in filas:
        ws.append([fila.get(h, "") for h in HEADERS])
    wb.save(destino)
    wb.close()
    return destino


def main() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
    )
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--plantilla",
        type=Path,
        default=Path(
            r"c:\Users\Diego\Downloads\cmct_g_20260805095021_43255_167_0_18775_0.xlsx"
        ),
    )
    parser.add_argument(
        "--ruts",
        type=str,
        default=(
            "11860305,14182089,15878960,16571229,16904091,17395456,"
            "19601890,19961115,20550934,20620627,21746929"
        ),
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=None,
    )
    args = parser.parse_args()
    ruts = [x.strip() for x in args.ruts.split(",") if x.strip()]
    out = args.out or (
        DOWNLOAD_DIR
        / "staging_trabajadores"
        / date.today().strftime("%Y_%m_%d")
        / "plantilla_ingreso_trabajadores.xlsx"
    )

    logger.info("Plantilla base: %s", args.plantilla)
    logger.info("RUTs: %s", ruts)
    filas, excluidos = construir_filas(plantilla=args.plantilla, ruts=ruts)
    for e in excluidos:
        logger.warning("Excluido: %s", e)
    logger.info("Filas a escribir: %s", len(filas))
    path = escribir_plantilla(args.plantilla, filas, out)
    logger.info("Escrito: %s", path)
    for f in filas:
        logger.info(
            "  %s | %s %s %s | faena=%s servicio=%s afp=%s inicio=%s ingreso_faena=%s",
            f["rut"],
            f["nombres"],
            f["apellido_paterno"],
            f["apellido_materno"],
            f["faena_id"],
            f["servicio_id"],
            f["afp_id"],
            f["fecha_inicio_contrato"],
            f["fecha_ingreso_faena"],
        )


if __name__ == "__main__":
    main()
